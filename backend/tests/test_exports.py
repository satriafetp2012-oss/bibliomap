import csv
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bibliomap.api import workspaces
from bibliomap.api.analysis import MapNode, ScienceMap
from bibliomap.exports import maps
from bibliomap.main import app
from bibliomap.workspace.pipeline import StateCache
from bibliomap.workspace.store import WorkspaceStore


def node(id_: str, label: str, cluster: int, **over) -> MapNode:
    base = dict(x=0.5, y=-0.25, weight=4, link_strength=7.5, documents=4, citations=12, year=2021.0)
    return MapNode(id=id_, label=label, cluster=cluster, **(base | over))


@pytest.fixture
def science_map() -> ScienceMap:
    return ScienceMap(
        type="co_occurrence",
        unit="keyword",
        normalization="association",
        nodes=[
            node("helmet", "helmet use", 1),
            node("motorcycle", 'motor"cycle', 1, x=-0.5, citations=3, year=None),
            node("injury", "injury", 2, y=0.75, weight=2),
        ],
        edges=[
            {"source": "helmet", "target": "motorcycle", "weight": 0.25, "raw_weight": 5},
            {"source": "helmet", "target": "injury", "weight": 0.125, "raw_weight": 2},
        ],
        clusters=2,
        modularity=0.42,
        iterations=300,
    )


# --- file formats ---------------------------------------------------------


def test_vosviewer_map_file(science_map):
    text = maps.vosviewer_map(science_map)

    assert text.startswith("﻿")
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿")), delimiter="\t"))
    assert rows[0] == [
        "id",
        "label",
        "x",
        "y",
        "cluster",
        "weight<Documents>",
        "weight<Total link strength>",
        "weight<Citations>",
        "score<Avg. pub. year>",
    ]
    # Ids are consecutive from 1, as VOSviewer requires.
    assert [r[0] for r in rows[1:]] == ["1", "2", "3"]
    assert rows[1][1] == "helmet use"
    assert rows[2][1] == 'motor"cycle'  # quoting survives the round trip
    assert rows[1][5] == "4" and rows[1][7] == "12"
    assert rows[2][8] == ""  # unknown year stays blank, not 0


def test_vosviewer_network_file_uses_the_same_ids(science_map):
    text = maps.vosviewer_network(science_map)
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿")), delimiter="\t"))
    assert rows == [["1", "2", "0.25"], ["1", "3", "0.125"]]


def test_cluster_colors_file(science_map):
    rows = list(
        csv.reader(io.StringIO(maps.cluster_colors(science_map).lstrip("﻿")), delimiter="\t")
    )
    assert rows[0] == ["cluster", "red", "green", "blue"]
    assert len(rows) == 3  # header + one row per cluster
    assert rows[1] == ["1", "37", "99", "235"]


def test_pajek_files(science_map):
    net = maps.pajek_network(science_map).splitlines()
    assert net[0] == "*Vertices 3"
    assert net[1] == '1 "helmet use"'
    # A quote inside a label would break Pajek's parser.
    assert net[2] == '2 "motor\'cycle"'
    assert net[4] == "*Edges"
    assert net[5] == "1 2 0.25"

    clu = maps.pajek_clusters(science_map).splitlines()
    assert clu == ["*Vertices 3", "1", "1", "2"]


def test_graphml_is_valid_xml_with_attributes(science_map):
    from xml.etree import ElementTree as ET

    root = ET.fromstring(maps.graphml(science_map))
    ns = {"g": "http://graphml.graphdrawing.org/xmlns"}
    nodes = root.findall(".//g:node", ns)
    edges = root.findall(".//g:edge", ns)
    assert len(nodes) == 3 and len(edges) == 2
    values = {d.get("key"): d.text for d in nodes[0].findall("g:data", ns)}
    assert values["label"] == "helmet use"
    assert values["cluster"] == "1"


def test_nodes_and_edges_csv(science_map):
    rows = list(csv.DictReader(io.StringIO(maps.nodes_csv(science_map))))
    assert rows[0]["label"] == "helmet use" and rows[0]["cluster"] == "1"
    assert rows[1]["year"] == ""

    edges = list(csv.DictReader(io.StringIO(maps.edges_csv(science_map))))
    # Edges carry labels, not internal ids, so the file is readable alone.
    assert edges[0]["source"] == "helmet use" and edges[0]["target"] == 'motor"cycle'
    assert edges[0]["raw_weight"] == "5"


# --- endpoints ------------------------------------------------------------


@pytest.fixture
def client(tmp_path):
    store = WorkspaceStore(tmp_path / "exports.sqlite3")
    cache = StateCache(store)
    app.dependency_overrides[workspaces.get_store] = lambda: store
    app.dependency_overrides[workspaces.get_cache] = lambda: cache
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def workspace(client, fixture_bytes):
    ws = client.post("/api/workspaces", json={"name": "Helmet review 2026"}).json()
    for name in ("wos_sample.txt", "scopus_sample.csv", "pubmed_sample.nbib"):
        client.post(
            f"/api/workspaces/{ws['id']}/imports/file", files={"file": (name, fixture_bytes(name))}
        )
    return ws["id"]


MAP_BODY = {"type": "co_authorship", "unit": "author", "min_occurrences": 1}


def test_export_map_vosviewer_zip(client, workspace):
    response = client.post(
        f"/api/workspaces/{workspace}/exports/map",
        params={"format": "vosviewer"},
        json=MAP_BODY,
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert "helmet-review-2026-co_authorship" in response.headers["content-disposition"]

    archive = zipfile.ZipFile(io.BytesIO(response.content))
    names = archive.namelist()
    assert any(n.endswith("-map.txt") for n in names)
    assert any(n.endswith("-network.txt") for n in names)
    assert "cluster_colors.txt" in names
    map_text = archive.read(next(n for n in names if n.endswith("-map.txt"))).decode("utf-8-sig")
    assert map_text.startswith("id\tlabel\tx\ty\tcluster")


@pytest.mark.parametrize(
    ("fmt", "media", "suffix"),
    [
        ("graphml", "application/xml", ".graphml"),
        ("pajek", "application/zip", "-pajek.zip"),
        ("csv", "application/zip", "-csv.zip"),
    ],
)
def test_export_map_other_formats(client, workspace, fmt, media, suffix):
    response = client.post(
        f"/api/workspaces/{workspace}/exports/map", params={"format": fmt}, json=MAP_BODY
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == media
    assert suffix in response.headers["content-disposition"]


def test_export_records_csv(client, workspace):
    response = client.get(f"/api/workspaces/{workspace}/exports/records", params={"format": "csv"})

    assert response.status_code == 200
    assert response.text.startswith("﻿")  # Excel needs the BOM
    rows = list(csv.DictReader(io.StringIO(response.text.lstrip("﻿"))))
    assert rows
    assert {"title", "authors", "year", "doi", "imports"} <= rows[0].keys()
    assert rows[0]["title"]
    assert all(r["exclusion_reason"] == "" for r in rows)


def test_export_records_csv_can_include_excluded(client, workspace):
    client.put(f"/api/workspaces/{workspace}/filters", json={"document_types": ["article"]})
    url = f"/api/workspaces/{workspace}/exports/records"
    kept = client.get(url, params={"format": "csv"}).text
    both = client.get(url, params={"format": "csv", "include_excluded": True}).text
    assert len(both) > len(kept)
    assert "document_type" in both


def test_export_workbook(client, workspace):
    response = client.get(f"/api/workspaces/{workspace}/exports/records")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert "-laporan.xlsx" in response.headers["content-disposition"]

    book = load_workbook(io.BytesIO(response.content))
    assert "Ringkasan" in book.sheetnames
    assert "Rekaman" in book.sheetnames
    assert "PRISMA" in book.sheetnames
    summary = {row[0]: row[1] for row in book["Ringkasan"].iter_rows(min_row=2, values_only=True)}
    assert summary["Workspace"] == "Helmet review 2026"
    assert summary["Dokumen"] >= 1
    records = book["Rekaman"]
    assert records.max_row >= 2
    assert records.freeze_panes == "A2"


def test_export_unknown_workspace_returns_404(client):
    assert client.get("/api/workspaces/nope/exports/records").status_code == 404
    assert client.post("/api/workspaces/nope/exports/map", json=MAP_BODY).status_code == 404
